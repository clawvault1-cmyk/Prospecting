"""Excel mail-merge workbook and a Word document with one prospect per page."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.config_loader import export_columns


def export_status(value: str | None) -> str:
    if (value or "").strip() == "verified":
        return "verified"
    return "needs review"


def build_export_records(rows: list[dict]) -> list[dict]:
    records = []
    for row in rows:
        first = row.get("first_name") or ""
        last = row.get("last_name") or ""
        records.append(
            {
                "first_name": first,
                "last_name": last,
                "full_name": " ".join(part for part in (first, last) if part).strip(),
                "email": row.get("email") or "",
                "title": row.get("title") or "",
                "company": row.get("company") or "",
                "subject": row.get("subject") or "",
                "body": row.get("body") or "",
                "followup_subject": row.get("followup_subject") or "",
                "followup_body": row.get("followup_body") or "",
                "trigger_source": row.get("trigger_source_url") or "",
                "status": export_status(row.get("email_status") or row.get("status")),
                "proof_point": row.get("proof_point_id") or "",
                "city": row.get("facility_city") or "",
            }
        )
    return records


def workbook_bytes(records: list[dict], claims_by_company: dict[str, list[dict]] | None = None) -> bytes:
    columns = export_columns()
    book = Workbook()
    sheet = book.active
    sheet.title = "Mail merge"
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(1, index, column["header"])
        cell.font = header_font
        cell.alignment = wrap
    for row_index, record in enumerate(records, start=2):
        for col_index, column in enumerate(columns, start=1):
            value = record.get(column["field"], "")
            cell = sheet.cell(row_index, col_index, value if value is not None else "")
            cell.alignment = wrap
        sheet.row_dimensions[row_index].height = 90
    widths = {
        "Body": 72,
        "FollowUpBody": 42,
        "Subject": 22,
        "FollowUpSubject": 22,
        "TriggerSource": 42,
        "Title": 28,
        "Company": 28,
        "Email": 28,
    }
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(column["header"], 16)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def document_bytes(records: list[dict], claims_for: list[list[dict]], sources_for: list[list[dict]]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    if not records:
        document.add_paragraph("No drafts to export.")
    for index, record in enumerate(records):
        if index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading(record.get("full_name") or record.get("company") or "Prospect", level=1)
        status = record.get("status") or "needs review"
        status_para = document.add_paragraph()
        run = status_para.add_run(f"Status: {status}")
        run.bold = True
        meta = document.add_paragraph()
        meta.add_run("To: ").bold = True
        meta.add_run(f"{record.get('full_name')} <{record.get('email')}>")
        role = document.add_paragraph()
        role.add_run("Role: ").bold = True
        role.add_run(f"{record.get('title')} — {record.get('company')}")
        subject = document.add_paragraph()
        subject.add_run("Subject: ").bold = True
        subject.add_run(record.get("subject") or "")
        document.add_paragraph("Body")
        _add_plain(document, record.get("body") or "(no body)")
        document.add_paragraph()
        follow_head = document.add_paragraph()
        follow_head.add_run("Day-3 follow-up").bold = True
        follow_subject = document.add_paragraph()
        follow_subject.add_run("Subject: ").bold = True
        follow_subject.add_run(record.get("followup_subject") or "")
        _add_plain(document, record.get("followup_body") or "(no follow-up)")
        trigger = document.add_paragraph()
        trigger.add_run("Trigger source: ").bold = True
        trigger.add_run(record.get("trigger_source") or "")

        document.add_heading("Sources", level=2)
        sources = sources_for[index] if index < len(sources_for) else []
        if not sources:
            document.add_paragraph("No sources stored.")
        for source in sources:
            paragraph = document.add_paragraph()
            paragraph.add_run(source.get("title") or source.get("url") or "Source").bold = True
            document.add_paragraph(source.get("url") or "")
            document.add_paragraph(source.get("snippet") or "")

        document.add_heading("Fact check", level=2)
        claims = claims_for[index] if index < len(claims_for) else []
        if not claims:
            document.add_paragraph("No claim table stored.")
        else:
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            headers = table.rows[0].cells
            headers[0].text = "Claim"
            headers[1].text = "Label"
            headers[2].text = "Support"
            for claim in claims:
                cells = table.add_row().cells
                cells[0].text = claim.get("claim") or ""
                cells[1].text = claim.get("label") or ""
                quote = claim.get("source_quote") or ""
                url = claim.get("source_url") or ""
                cells[2].text = (quote + (f"\n{url}" if url else "")).strip()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_plain(document: Document, text: str) -> None:
    for line in (text or "").split("\n"):
        document.add_paragraph(line)

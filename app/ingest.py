"""Read a ZoomInfo-style CSV or Excel export into plain row dicts."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from openpyxl import load_workbook

HEADER_MAP = {
    "first name": "first_name",
    "last name": "last_name",
    "job title": "title",
    "management level": "management_level",
    "job function": "job_function",
    "department": "department",
    "email address": "email",
    "person city": "person_city",
    "person state": "person_state",
    "company name": "company",
    "website": "website",
    "revenue (in 000s usd)": "revenue_000s",
    "revenue range (in usd)": "revenue_range",
    "employees": "employees",
    "employee range": "employee_range",
    "primary industry": "primary_industry",
    "primary sub-industry": "primary_sub_industry",
    "all industries": "all_industries",
    "all sub-industries": "all_sub_industries",
    "company city": "company_city",
    "company state": "company_state",
    "company country": "company_country",
    "number of locations": "locations",
}

# Street fields are recognized so they are dropped on purpose.
DROPPED = {"person street", "company street address", "full address"}


def rows_from_upload(filename: str, payload: bytes) -> list[dict]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return rows_from_csv_bytes(payload)
    if lower.endswith(".xlsx"):
        return rows_from_xlsx_bytes(payload)
    raise ValueError("Upload a .csv or .xlsx file.")


def rows_from_csv_bytes(payload: bytes) -> list[dict]:
    text = _decode(payload)
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("The CSV has no header row.")
    return [_map_row(row) for row in reader]


def rows_from_xlsx_bytes(payload: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            raise ValueError("The workbook is empty.")
        headers = [_cell_str(value) for value in header_row]
        mapped = []
        for values in iterator:
            raw = {headers[index]: _cell_str(value) for index, value in enumerate(values) if index < len(headers)}
            mapped.append(_map_row(raw))
        return mapped
    finally:
        workbook.close()


def rows_from_path(path: Path) -> list[dict]:
    return rows_from_upload(path.name, path.read_bytes())


def _decode(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _map_row(raw: dict) -> dict:
    mapped = {key: "" for key in set(HEADER_MAP.values())}
    for header, value in raw.items():
        if header is None:
            continue
        key = re.sub(r"\s+", " ", str(header).strip().lower())
        if key in DROPPED:
            continue
        field = HEADER_MAP.get(key)
        if not field:
            continue
        mapped[field] = _cell_str(value)
    return mapped

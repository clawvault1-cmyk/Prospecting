from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app import db
from app.config_loader import export_columns
from app.exporters import build_export_records, document_bytes, export_status, workbook_bytes
from app.ingest import rows_from_path
from app.pipeline import process_company
from app.scoring import score_rows

FIXTURE = Path(__file__).parent / "fixtures" / "fake_zoominfo.csv"


def test_status_labels():
    assert export_status("verified") == "verified"
    assert export_status("needs_review") == "needs review"
    assert export_status("") == "needs review"


def test_workbook_and_docx_from_fixture_run():
    batch_id = db.insert_batch("fake_zoominfo.csv", score_rows(rows_from_path(FIXTURE)))
    for company in db.list_companies(batch_id):
        if company["included"]:
            process_company(company["id"])

    rows = db.list_export_rows(batch_id)
    names = {row["company"] for row in rows}
    assert "Harborline Foods" not in names
    assert "Northline Distribution" in names
    by_company = {row["company"]: row for row in rows}
    assert by_company["Little Bolt Co"]["email_status"] == "verified"
    assert by_company["Northline Distribution"]["email_status"] == "verified"
    records = build_export_records(rows)
    north = next(record for record in records if record["company"] == "Northline Distribution")
    assert north["status"] == "verified"
    assert "\n" in north["body"]
    assert north["full_name"] == "Maya Chen"

    payload = workbook_bytes(records)
    book = load_workbook(BytesIO(payload))
    sheet = book.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == [column["header"] for column in export_columns()]
    body_index = headers.index("Body") + 1
    status_index = headers.index("Status") + 1
    values = [sheet.cell(row, body_index).value or "" for row in range(2, sheet.max_row + 1)]
    assert any("\n" in value for value in values)
    statuses = {sheet.cell(row, status_index).value for row in range(2, sheet.max_row + 1)}
    assert statuses == {"verified"}

    claims = []
    sources = []
    for row in rows:
        company = db.get_company(row["company_id"])
        claims.append(company["claims"])
        sources.append(company["sources"])
    document = Document(BytesIO(document_bytes(records, claims, sources)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Maya Chen" in text
    assert "Status: verified" in text
    assert "Fact check" in "".join(paragraph.text for paragraph in document.paragraphs) or document.tables
    xml = document.element.xml
    assert 'w:type="page"' in xml
    assert xml.count('w:type="page"') == len(records) - 1
